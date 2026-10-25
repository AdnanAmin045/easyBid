"""Read a recipient list: pasted addresses, free text, a CSV/TSV sheet, or an Excel file.

A header row maps columns to the recipient: email (required), name, company, subject and body. Every other column
is kept and becomes a template variable. When a row has its own body, it is sent as written.

Without a header, the columns are told apart by what they contain, in any order: the one full of addresses, then the
longest text is the message and a shorter one the subject. An email may also start with a line "Subject: ...."
"""

import csv
import io
import itertools
import re
from dataclasses import dataclass, field

from email_validator import EmailNotValidError, validate_email

from app.jobs.extract import _EMAIL, deobfuscate, find_emails
from app.outreach.render import key

ALIASES = {
    "email": {"email", "e_mail", "email_address", "mail", "address", "recipient", "to", "receiver", "recipient_email", "to_email", "email_id", "mail_id"},
    "name": {"name", "full_name", "fullname", "contact", "contact_name", "person"},
    "first_name": {"first_name", "firstname", "first", "given_name"},
    "last_name": {"last_name", "lastname", "last", "surname", "family_name"},
    "company": {"company", "company_name", "organization", "organisation", "employer", "business"},
    "subject": {"subject", "email_subject"},
    "body": {"body", "email_body", "message", "content", "mail_body", "full_email", "email_text", "mail_text", "text", "draft", "email_message"},
}
DELIMITERS = (",", "\t", ";", "|")
SUBJECT_LINE = re.compile(r"\A\s*subject\s*:[ \t]*([^\n]+)\n+", re.I)
NAMED_ADDRESS = re.compile(r'^\s*"?([^"<>]*?)"?\s*<\s*([^<>\s]+@[^<>\s]+)\s*>\s*$')


@dataclass
class Row:
    email: str
    name: str = ""
    company: str = ""
    subject: str = ""
    body: str = ""
    fields: dict = field(default_factory=dict)
    line: int = 0


@dataclass
class Rejected:
    line: int
    value: str
    reason: str


@dataclass
class Sheet:
    rows: list[Row] = field(default_factory=list)
    rejected: list[Rejected] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)  # extra columns found, usable as {{variables}}
    has_header: bool = False


def clean_address(value: str) -> tuple[str, str] | None:
    """(email, name) from 'a@b.com' or 'Ann <a@b.com>', or None when it is not a valid address."""
    name = ""
    named = NAMED_ADDRESS.match(value)
    if named:
        name, value = named.group(1).strip(), named.group(2)
    try:
        email = validate_email(value.strip().strip("<>;,"), check_deliverability=False).normalized.lower()
    except EmailNotValidError:
        return None
    return email, name


def _header_map(cells: list[str]) -> dict[int, str] | None:
    """Column index -> role (or the column's own key) when this row is a header with an email column."""
    roles: dict[int, str] = {}
    for i, cell in enumerate(cells):
        name = key(cell)
        if not name:
            continue
        role = next((r for r, names in ALIASES.items() if name in names), None)
        roles[i] = role or name
    return roles if "email" in roles.values() else None


def _header(text: str) -> tuple[str, dict[int, str]] | None:
    first = next((line for line in text.splitlines() if line.strip()), "")
    for delimiter in DELIMITERS:
        mapping = _header_map(next(csv.reader([first], delimiter=delimiter), []))
        if mapping:
            return delimiter, mapping
    return None


def _only_addresses(cell: str) -> bool:
    """True for 'bob@x.io; carol [at] y [dot] io': addresses and separators only, so a list and not an email text."""
    rest = _EMAIL.sub("", deobfuscate(cell))
    return not re.search(r"[A-Za-z]", rest)


def infer_columns(rows: list[list[str]]) -> dict[int, str] | None:
    """Work out which column is the address, the subject and the message from what the cells contain.

    The address column is the one full of valid addresses. Of the others, the one with the longest text is the
    message and a shorter one is the subject. None when the sheet does not look like that, e.g. a plain list
    of addresses spread over several columns.
    """
    sample = [r for r in rows if any(c.strip() for c in r)][:20]
    width = max((len(r) for r in sample), default=0)
    if width < 2:
        return None

    def column(i: int) -> list[str]:
        return [r[i].strip() for r in sample if i < len(r) and r[i].strip()]

    cols = {i: column(i) for i in range(width)}
    cols = {i: cells for i, cells in cols.items() if cells}
    address = [i for i, cells in cols.items() if sum(bool(clean_address(c)) for c in cells) >= 0.6 * len(cells)]
    others = [i for i in cols if i not in address]
    if len(address) != 1 or not others:
        return None
    if any(sum(_only_addresses(c) for c in cols[i]) > len(cols[i]) / 2 for i in others):
        return None  # "a@x.io, b@y.io": more addresses, not an email text
    others.sort(key=lambda i: sum(len(c) for c in cols[i]) / len(cols[i]))
    mapping = {address[0]: "email", others[-1]: "body"}
    if len(others) >= 2:
        mapping[others[0]] = "subject"
    return mapping


def _positional(text: str) -> tuple[str, dict[int, str]] | None:
    """(delimiter, columns) for a sheet without a header row."""
    best = None
    for delimiter in DELIMITERS:
        rows = list(itertools.islice(csv.reader(io.StringIO(text), delimiter=delimiter), 40))
        mapping = infer_columns(rows)
        if mapping and (best is None or len(mapping) > len(best[1])):
            best = (delimiter, mapping)
    return best


def _from_rows(rows, mapping: dict[int, str], has_header: bool) -> Sheet:
    """rows: (line number, cells) pairs, without the header row."""
    sheet = Sheet(has_header=has_header)
    seen: set[str] = set()
    extra = [role for role in mapping.values() if role not in ALIASES]
    sheet.columns = list(dict.fromkeys(extra))
    for line, cells in rows:
        if not any(c.strip() for c in cells):
            continue
        values: dict[str, str] = {}
        for i, role in mapping.items():
            values.setdefault(role, cells[i].strip() if i < len(cells) else "")
        raw = values.get("email", "")
        address = clean_address(raw) if raw else None
        if not address:
            sheet.rejected.append(Rejected(line, raw[:200], "Not a valid email address" if raw else "No email in this row"))
            continue
        email, embedded_name = address
        if email in seen:
            sheet.rejected.append(Rejected(line, email, "Duplicate in this list"))
            continue
        seen.add(email)
        name = values.get("name") or " ".join(p for p in (values.get("first_name"), values.get("last_name")) if p) or embedded_name
        subject = " ".join(values.get("subject", "").split())[:300]
        body = values.get("body", "").replace("\r\n", "\n").replace("\r", "\n").strip()
        found = SUBJECT_LINE.match(body + "\n")
        if found and not subject:
            subject, body = " ".join(found.group(1).split())[:300], body[found.end() - 1 :].strip()
        sheet.rows.append(
            Row(
                email=email,
                name=name[:200],
                company=values.get("company", "")[:200],
                subject=subject,
                body=body,
                fields={c: values[c] for c in sheet.columns if values.get(c)},
                line=line,
            )
        )
    return sheet


def _csv_rows(text: str, delimiter: str, skip_first: bool):
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    if skip_first:
        next(reader, None)
    for cells in reader:
        yield reader.line_num, cells


def _from_free_text(text: str) -> Sheet:
    sheet = Sheet()
    seen: set[str] = set()
    for number, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        named = NAMED_ADDRESS.match(line)
        candidates = [line] if named else find_emails(line)
        found = [address for address in map(clean_address, candidates) if address]
        if not found:
            sheet.rejected.append(Rejected(number, line.strip()[:200], "No email address found"))
            continue
        for email, name in found:
            if email in seen:
                sheet.rejected.append(Rejected(number, email, "Duplicate in this list"))
                continue
            seen.add(email)
            sheet.rows.append(Row(email=email, name=name[:200], line=number))
    return sheet


def parse(text: str) -> Sheet:
    text = text.lstrip("﻿")
    header = _header(text)
    if header:
        delimiter, mapping = header
        return _from_rows(_csv_rows(text, delimiter, True), mapping, True)
    positional = _positional(text)
    if positional:
        return _from_rows(_csv_rows(text, positional[0], False), positional[1], False)
    return _from_free_text(text)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(data: bytes, limit: int = 20000) -> Sheet:
    """The first worksheet of an Excel file, read like a sheet pasted from it."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError, Exception) as e:  # a corrupt zip raises many kinds
        raise ValueError("This is not a readable Excel (.xlsx) file") from e
    sheet = workbook.worksheets[0]
    rows: list[tuple[int, list[str]]] = []
    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if number > limit:
            break
        rows.append((number, [_cell(v) for v in row]))
    workbook.close()
    first = next((cells for _, cells in rows if any(c.strip() for c in cells)), [])
    mapping = _header_map(first)
    if mapping:
        return _from_rows(rows[1:], mapping, True)
    return _from_rows(rows, infer_columns([cells for _, cells in rows]) or {0: "email", 1: "body"}, False)
